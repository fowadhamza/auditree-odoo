from datetime import date
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

output_path = "/home/fowadhamza/Projects/Auditree/Shaik_Danyal_Leave_Tally_2026-08-27.xlsx"
cutoff = date(2026, 8, 27)
summary = [("Casual Leave", 10.0, 6.0, 4.0), ("Sick Leave", 8.0, 6.5, 1.5), ("Comp-off", 1.5, 1.5, 0.0), ("Earned Leave balance b/f", 5.0, 5.0, 0.0)]
entries = [(date(2026, 1, 9), "Sick Leave", 0.5), (date(2026, 1, 12), "Earned Leave balance b/f", 1.0), (date(2026, 2, 11), "Earned Leave balance b/f", 1.0), (date(2026, 3, 20), "Earned Leave balance b/f", 1.0), (date(2026, 4, 2), "Earned Leave balance b/f", 0.5), (date(2026, 4, 15), "Earned Leave balance b/f", 0.5), (date(2026, 4, 16), "Casual Leave", 1.0), (date(2026, 5, 20), "Earned Leave balance b/f", 0.5), (date(2026, 5, 28), "Casual Leave", 2.0), (date(2026, 6, 5), "Casual Leave", 1.0), (date(2026, 6, 9), "Earned Leave balance b/f", 0.5), (date(2026, 6, 25), "Casual Leave", 1.0), (date(2026, 7, 2), "Sick Leave", 1.0), (date(2026, 7, 7), "Sick Leave", 1.0), (date(2026, 7, 13), "Sick Leave", 4.0), (date(2026, 7, 17), "Casual Leave", 1.0), (date(2026, 8, 7), "Comp-off", 1.0), (date(2026, 8, 10), "Comp-off", 0.5)]

workbook = Workbook()
summary_sheet = workbook.active
summary_sheet.title = "Balance Summary"
summary_sheet.append(["Employee", "Shaik Danyal"])
summary_sheet.append(["Production database", "AUDITREE_LIVE_NEW"])
summary_sheet.append(["As-of date", cutoff])
summary_sheet.append([])
summary_sheet.append(["Leave Type", "Allocated", "Taken", "Balance"])
for row in summary:
    summary_sheet.append(row)
summary_sheet.append(["Total", sum(row[1] for row in summary), sum(row[2] for row in summary), sum(row[3] for row in summary)])
summary_sheet.append([])
summary_sheet.append(["Note", "Taken values include validated 2026 leave dated on or before the as-of date."])

entries_sheet = workbook.create_sheet("Approved Leave Taken")
entries_sheet.append(["Employee", "Leave Date", "Leave Type", "Days Taken", "Status"])
for leave_date, leave_type, days in entries:
    entries_sheet.append(["Shaik Danyal", leave_date, leave_type, days, "Validated"])
entries_sheet.append([])
entries_sheet.append(["Note", "Only validated 2026 entries through 2026-08-27 are included in the summary."])

header_fill = PatternFill("solid", fgColor="1F4E78")
for sheet in workbook.worksheets:
    for cell in sheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
    sheet.freeze_panes = "A2"
    for row in sheet.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top")
    for column_cells in sheet.columns:
        width = max(len(str(cell.value or "")) for cell in column_cells) + 2
        sheet.column_dimensions[get_column_letter(column_cells[0].column)].width = min(width, 34)

for cell in summary_sheet[3]:
    cell.number_format = "yyyy-mm-dd"
for row in entries_sheet.iter_rows(min_row=2, max_row=entries_sheet.max_row):
    row[1].number_format = "yyyy-mm-dd"

workbook.save(output_path)
print(output_path)
